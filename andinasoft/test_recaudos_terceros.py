import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
from django.test import SimpleTestCase

from andinasoft.recaudos_terceros_service import (
    agrupar_por_tercero, datos_basicos_terceros, excel_recaudos_terceros, orden_recibo, repartir,
)

D = Decimal


def _recibo(num, adj, valor, fecha=date(2026, 10, 1), tercero='1', forma='BANCO'):
    return SimpleNamespace(numrecibo=num, idadjudicacion=adj, idtercero=tercero, valor=D(valor),
                           fecha=fecha, formapago=forma)


class RepartirTests(SimpleTestCase):
    def test_un_titular(self):
        self.assertEqual(repartir(D('100'), 1), [D('100')])

    def test_residuo_al_primero(self):
        partes = repartir(D('100'), 3)
        self.assertEqual(partes, [D('33.34'), D('33.33'), D('33.33')])
        self.assertEqual(sum(partes), D('100'))


class AgruparPorTerceroTests(SimpleTestCase):
    def test_reparte_entre_titulares_y_cuadra(self):
        recibos = [_recibo('M1', 'ADJ1', '1000')]
        detalle = {'M1': {'capital': D('700'), 'interescte': D('200'), 'interesmora': D('50'), 'saldo_favor': D('0')}}
        informe = agrupar_por_tercero(
            recibos, detalle, {'ADJ1': ['1', '2']}, {'1': 'ana', '2': 'luis'},
            {'BANCO': '11100501'}, {'capital': '13050501', 'saldo_favor': '28050501'},
        )
        self.assertEqual([g['nombre'] for g in informe['terceros']], ['ANA', 'LUIS'])
        ana = informe['terceros'][0]['movimientos'][0]
        self.assertEqual(ana['capital'], D('350'))
        self.assertEqual(ana['sin_aplicar'], D('25'))
        self.assertEqual(ana['participacion'], D('50.00'))
        self.assertEqual(informe['totales']['total'], D('1000'))
        self.assertEqual(informe['total_debito'], informe['total_credito'])
        self.assertEqual(informe['resumen'][0]['cuenta'], '11100501')

    def test_agrupa_varios_recibos_del_mismo_tercero(self):
        recibos = [
            _recibo('M2', 'ADJ1', '50', fecha=date(2026, 10, 5)),
            _recibo('M1', 'ADJ1', '30', fecha=date(2026, 10, 2)),
        ]
        detalle = {r.numrecibo: {'capital': r.valor} for r in recibos}
        informe = agrupar_por_tercero(recibos, detalle, {'ADJ1': ['1']}, {'1': 'Ana'}, {})
        grupo = informe['terceros'][0]
        self.assertEqual([m['recibo'] for m in grupo['movimientos']], ['M1', 'M2'])
        self.assertEqual(grupo['totales']['capital'], D('80'))
        self.assertEqual(informe['num_recibos'], 2)

    def test_sin_adjudicacion_usa_tercero_del_recibo(self):
        informe = agrupar_por_tercero([_recibo('M1', 'ADJ9', '10', tercero='77')], {}, {}, {}, {})
        grupo = informe['terceros'][0]
        self.assertEqual(grupo['nit'], '77')
        self.assertEqual(grupo['totales']['sin_aplicar'], D('10'))

    def test_vista_por_recibo_en_orden_con_sus_titulares(self):
        recibos = [_recibo('10', 'ADJ1', '100'), _recibo('9', 'ADJ2', '40')]
        detalle = {'10': {'capital': D('100')}, '9': {'capital': D('40')}}
        informe = agrupar_por_tercero(
            recibos, detalle, {'ADJ1': ['1', '2'], 'ADJ2': ['3']}, {'1': 'Ana', '2': 'Luis', '3': 'Eva'}, {},
        )
        self.assertEqual([r['recibo'] for r in informe['recibos']], ['9', '10'])
        diez = informe['recibos'][1]
        self.assertEqual([(m['nit'], m['nombre'], m['capital']) for m in diez['lineas']],
                         [('1', 'ANA', D('50')), ('2', 'LUIS', D('50'))])
        self.assertEqual(diez['totales']['total'], D('100'))

    def test_orden_recibo_natural(self):
        self.assertEqual(sorted(['M10', '10', 'M9', '9', '100'], key=orden_recibo), ['9', '10', '100', 'M9', 'M10'])

    def test_excel_lineas_con_tercero_y_hoja_de_terceros(self):
        informe = agrupar_por_tercero([_recibo('M1', 'ADJ1', '10')], {}, {'ADJ1': ['1', '2']},
                                      {'1': 'Ana', '2': 'Luis'}, {})
        informe['datos_terceros'] = datos_basicos_terceros(informe['terceros'], {
            '1': {'tipo_doc': '13', 'domicilio': 'Calle 1', 'ciudad': 'MONTERIA', 'email': 'ana@x.co',
                  'celular1': '300'},
        })
        for vista, activa in (('tercero', 'Por tercero'), ('recibo', 'Por recibo')):
            contenido = excel_recaudos_terceros(informe, 'Oasis', date(2026, 10, 1), date(2026, 10, 31), vista)
            book = openpyxl.load_workbook(io.BytesIO(contenido))
            self.assertEqual(book.active.title, activa)
        hoja = book['Por recibo']
        self.assertEqual([hoja.cell(4, c).value for c in (5, 6)], ['Id tercero', 'Nombre tercero'])
        self.assertEqual([(hoja.cell(r, 2).value, hoja.cell(r, 5).value, hoja.cell(r, 6).value) for r in (5, 6)],
                         [('M1', '1', 'ANA'), ('M1', '2', 'LUIS')])
        terceros = book['Terceros']
        self.assertEqual([c.value for c in terceros[2]],
                         ['Cédula de ciudadanía', '1', 'ANA', 'Calle 1', 'MONTERIA', None, None, 'ana@x.co', '300', None])
        self.assertEqual(terceros.cell(3, 3).value, 'LUIS')
