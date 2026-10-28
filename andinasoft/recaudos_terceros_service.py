"""
Informe contable de recaudos por tercero (auxiliar por tercero).

Reparte cada recibo entre los titulares de la adjudicación con la misma regla que
la interfaz de recibos a SIIGO (`Recaudos_general.interfaz_contabilidad`): partes
iguales por titular y crédito por concepto (capital, interés corriente, interés de
mora y saldo a favor). Lo que el recibo no tiene aplicado en `recaudos` queda en
"sin aplicar" para que el total de cada recibo siga cuadrando con su valor.
"""
import io
import re
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from andinasoft.models import clientes
from andinasoft.shared_models import Adjudicacion, Recaudos, Recaudos_general, consecutivos, formas_pago

CONCEPTOS = ('capital', 'interescte', 'interesmora', 'saldo_favor', 'sin_aplicar')
_CENT = Decimal('0.01')


def _dec(x):
    if x is None:
        return Decimal('0')
    return x if isinstance(x, Decimal) else Decimal(str(x))


def _ceros():
    return {c: Decimal('0') for c in CONCEPTOS + ('total',)}


def repartir(valor, partes):
    """Divide `valor` en `partes` iguales a 2 decimales; el residuo va a la primera parte."""
    valor = _dec(valor)
    if partes <= 1:
        return [valor]
    cuota = (valor / partes).quantize(_CENT, rounding=ROUND_HALF_UP)
    return [valor - cuota * (partes - 1)] + [cuota] * (partes - 1)


def orden_recibo(numrecibo):
    """Orden natural del consecutivo: '9' < '10' y 'M9' < 'M10'."""
    return [(0, int(p), '') if p.isdigit() else (1, 0, p) for p in re.split(r'(\d+)', numrecibo or '') if p]


def agrupar_por_tercero(recibos, detalle, titulares_adj, nombres, cuentas_fp, cuentas_rc=None):
    """
    recibos: objetos con numrecibo, fecha, idadjudicacion, idtercero, valor, formapago.
    detalle: {numrecibo: {'capital', 'interescte', 'interesmora', 'saldo_favor'}}.
    titulares_adj: {idadjudicacion: [idtercero, ...]} en orden titular 1..4.
    nombres: {idtercero: nombrecompleto}.
    cuentas_fp: {formapago: cuenta_contable}.
    cuentas_rc: {'capital': cuenta, 'saldo_favor': cuenta} del consecutivo RC.
    """
    cuentas_rc = cuentas_rc or {}
    terceros = {}
    recibos_out = []
    debitos = {}
    totales = _ceros()

    for r in recibos:
        valor = _dec(r.valor)
        det = detalle.get(r.numrecibo, {})
        montos = {c: _dec(det.get(c)) for c in CONCEPTOS[:-1]}
        montos['sin_aplicar'] = valor - sum(montos.values())

        ids = [t for t in titulares_adj.get(r.idadjudicacion, []) if t] or [r.idtercero or '']
        partes = {c: repartir(montos[c], len(ids)) for c in CONCEPTOS}
        participacion = (Decimal('100') / len(ids)).quantize(_CENT, rounding=ROUND_HALF_UP)
        lineas = []

        for i, nit in enumerate(ids):
            grupo = terceros.get(nit)
            if grupo is None:
                grupo = terceros[nit] = {
                    'nit': nit,
                    'nombre': (nombres.get(nit) or ('TERCERO NO REGISTRADO' if nit else 'SIN TERCERO')).upper(),
                    'movimientos': [],
                    'totales': _ceros(),
                }
            mov = {c: partes[c][i] for c in CONCEPTOS}
            mov['total'] = sum(mov.values())
            mov.update({
                'fecha': r.fecha,
                'recibo': r.numrecibo,
                'contrato': r.idadjudicacion,
                'forma_pago': r.formapago or '',
                'participacion': participacion,
                'nit': nit,
                'nombre': grupo['nombre'],
            })
            grupo['movimientos'].append(mov)
            lineas.append(mov)
            for c in CONCEPTOS + ('total',):
                grupo['totales'][c] += mov[c]
                totales[c] += mov[c]

        recibos_out.append({
            'recibo': r.numrecibo,
            'fecha': r.fecha,
            'contrato': r.idadjudicacion,
            'forma_pago': r.formapago or '',
            'lineas': lineas,
            'totales': dict(montos, total=valor),
        })

        forma = r.formapago or ''
        deb = debitos.setdefault(forma, {
            'concepto': forma or 'SIN FORMA DE PAGO',
            'cuenta': cuentas_fp.get(forma, ''),
            'recibos': 0,
            'debito': Decimal('0'),
            'credito': Decimal('0'),
        })
        deb['recibos'] += 1
        deb['debito'] += valor

    for grupo in terceros.values():
        grupo['movimientos'].sort(key=lambda m: (m['fecha'], m['recibo']))

    creditos = [
        {'concepto': 'Capital', 'cuenta': cuentas_rc.get('capital', ''), 'credito': totales['capital']},
        {'concepto': 'Interés corriente', 'cuenta': cuentas_rc.get('capital', ''), 'credito': totales['interescte']},
        {'concepto': 'Interés de mora', 'cuenta': cuentas_rc.get('capital', ''), 'credito': totales['interesmora']},
        {'concepto': 'Saldo a favor', 'cuenta': cuentas_rc.get('saldo_favor', ''), 'credito': totales['saldo_favor']},
        {'concepto': 'Sin aplicar', 'cuenta': '', 'credito': totales['sin_aplicar']},
    ]
    resumen = sorted(debitos.values(), key=lambda d: d['concepto']) + [
        {'concepto': c['concepto'], 'cuenta': c['cuenta'], 'recibos': None,
         'debito': Decimal('0'), 'credito': c['credito']}
        for c in creditos if c['credito']
    ]
    total_debito = sum(d['debito'] for d in debitos.values())

    return {
        'terceros': sorted(terceros.values(), key=lambda g: (g['nombre'], g['nit'])),
        'recibos': sorted(recibos_out, key=lambda r: orden_recibo(r['recibo'])),
        'totales': totales,
        'resumen': resumen,
        'total_debito': total_debito,
        'total_credito': totales['total'],
        'num_recibos': sum(d['recibos'] for d in debitos.values()),
    }


def informe_recaudos_terceros(proyecto, desde, hasta):
    """Arma el informe para un proyecto entre dos fechas (inclusive). Excluye recibos anulados (valor 0)."""
    recibos_qs = Recaudos_general.objects.using(proyecto).filter(
        fecha__range=(desde, hasta), valor__gt=0,
    ).order_by('fecha', 'numrecibo')
    recibos = list(recibos_qs)

    detalle = defaultdict(lambda: dict.fromkeys(CONCEPTOS[:-1], Decimal('0')))
    lineas = Recaudos.objects.using(proyecto).filter(
        recibo__in=recibos_qs.values('numrecibo'),
    ).values_list('recibo', 'idcta', 'capital', 'interescte', 'interesmora')
    for recibo, idcta, capital, intcte, intmora in lineas:
        d = detalle[recibo]
        if (idcta or '').startswith('SF'):
            d['saldo_favor'] += _dec(capital)
        else:
            d['capital'] += _dec(capital)
            d['interescte'] += _dec(intcte)
            d['interesmora'] += _dec(intmora)

    adjs = {r.idadjudicacion for r in recibos if r.idadjudicacion}
    titulares_adj = {
        pk: [t for t in (t1, t2, t3, t4) if t]
        for pk, t1, t2, t3, t4 in Adjudicacion.objects.using(proyecto).filter(pk__in=adjs).values_list(
            'pk', 'idtercero1', 'idtercero2', 'idtercero3', 'idtercero4',
        )
    }

    ids = {r.idtercero for r in recibos if r.idtercero}
    for lista in titulares_adj.values():
        ids.update(lista)
    datos = {
        c['pk']: c for c in clientes.objects.filter(pk__in=ids).values(
            'pk', 'tipo_doc', 'nombrecompleto', 'domicilio', 'ciudad', 'city__city_name',
            'estado__state_name', 'pais__country_name', 'email', 'celular1', 'telefono1',
        )
    }
    nombres = {pk: c['nombrecompleto'] for pk, c in datos.items()}

    cuentas_fp = dict(formas_pago.objects.using(proyecto).values_list('descripcion', 'cuenta_contable'))
    rc = consecutivos.objects.using(proyecto).filter(documento='RC').first()
    cuentas_rc = {}
    if rc:
        cuentas_rc = {
            'capital': rc.cuenta_capital or '',
            'saldo_favor': (rc.cuenta_aux1 or '').strip() or rc.cuenta_capital or '',
        }

    informe = agrupar_por_tercero(recibos, detalle, titulares_adj, nombres, cuentas_fp, cuentas_rc)
    informe['datos_terceros'] = datos_basicos_terceros(informe['terceros'], datos)
    return informe


def datos_basicos_terceros(terceros, datos):
    """Una fila por tercero del informe con su información de contacto de `clientes`."""
    tipos = dict(clientes.dc_types)
    filas = []
    for g in terceros:
        c = datos.get(g['nit']) or {}
        filas.append({
            'tipo_doc': tipos.get(c.get('tipo_doc'), c.get('tipo_doc') or ''),
            'nit': g['nit'],
            'nombre': g['nombre'],
            'direccion': (c.get('domicilio') or '').strip(),
            'ciudad': (c.get('city__city_name') or c.get('ciudad') or '').strip(),
            'departamento': c.get('estado__state_name') or '',
            'pais': c.get('pais__country_name') or '',
            'email': c.get('email') or '',
            'celular': c.get('celular1') or '',
            'telefono': c.get('telefono1') or '',
        })
    return filas


_MONTOS = (('Capital', 16), ('Int. corriente', 16), ('Int. mora', 16), ('Saldo a favor', 16),
           ('Sin aplicar', 16), ('Total', 18))
_COLS_LINEA = (('Fecha', 12), ('Recibo', 12), ('Contrato', 12), ('Forma de pago', 24), ('Id tercero', 16),
               ('Nombre tercero', 40), ('% Part.', 9)) + _MONTOS
_COLS_TERCEROS = (('Tipo documento', 24), ('Cédula / NIT', 16), ('Nombre', 40), ('Dirección', 40), ('Ciudad', 20),
                  ('Departamento', 18), ('País', 14), ('Email', 32), ('Celular', 16), ('Teléfono', 16))
_NUM = '#,##0.00;[Red]-#,##0.00'
_NEGRITA = Font(bold=True)
_ENCABEZADO = PatternFill('solid', fgColor='1F3864')
_SUBTOTAL = PatternFill('solid', fgColor='F2F2F2')
_LINEA = Side(style='thin', color='999999')


def _encabezados(sheet, row, columnas):
    for col, (nombre, w) in enumerate(columnas, start=1):
        celda = sheet.cell(row, col, nombre)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = _ENCABEZADO
        celda.alignment = Alignment(horizontal='center')
        sheet.column_dimensions[get_column_letter(col)].width = w


def _hoja_agrupada(sheet, titulo, subtitulo, grupos, totales):
    """
    Escribe un auxiliar en el que cada línea lleva recibo y tercero completos, con subtotal por grupo.
    grupos: [(movimientos, texto_subtotal, totales_grupo)].
    """
    sheet['A1'] = titulo
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A2'] = subtitulo
    ancho = len(_COLS_LINEA)
    inicio_montos = ancho - len(_MONTOS) + 1
    _encabezados(sheet, 4, _COLS_LINEA)
    sheet.freeze_panes = 'A5'
    row = 5

    def _montos(fila, datos):
        for col, c in enumerate(CONCEPTOS + ('total',), start=inicio_montos):
            sheet.cell(fila, col, float(datos[c])).number_format = _NUM

    for movimientos, texto_subtotal, totales_grupo in grupos:
        for m in movimientos:
            sheet.cell(row, 1, m['fecha']).number_format = 'DD/MM/YYYY'
            sheet.cell(row, 2, m['recibo'])
            sheet.cell(row, 3, m['contrato'])
            sheet.cell(row, 4, m['forma_pago'])
            sheet.cell(row, 5, m['nit'])
            sheet.cell(row, 6, m['nombre'])
            sheet.cell(row, 7, float(m['participacion']) / 100).number_format = '0.##%'
            _montos(row, m)
            row += 1
        sheet.cell(row, 1, texto_subtotal)
        _montos(row, totales_grupo)
        for col in range(1, ancho + 1):
            sheet.cell(row, col).fill = _SUBTOTAL
            sheet.cell(row, col).font = _NEGRITA
            sheet.cell(row, col).border = Border(top=_LINEA)
        row += 1

    row += 1
    sheet.cell(row, 1, 'TOTAL GENERAL')
    _montos(row, totales)
    for col in range(1, ancho + 1):
        sheet.cell(row, col).font = _NEGRITA
        sheet.cell(row, col).border = Border(top=_LINEA, bottom=Side(style='double'))


def excel_recaudos_terceros(informe, proyecto, desde, hasta, vista='tercero'):
    """Devuelve los bytes de un .xlsx con las vistas por tercero y por recibo, los terceros y el resumen contable."""
    book = openpyxl.Workbook()
    proyecto = str(proyecto).upper()
    subtitulo = (f'Desde {desde:%d/%m/%Y} hasta {hasta:%d/%m/%Y} · {informe["num_recibos"]} recibos · '
                 f'{len(informe["terceros"])} terceros · Valores repartidos entre titulares como en la interfaz contable')

    por_tercero = book.active
    por_tercero.title = 'Por tercero'
    _hoja_agrupada(
        por_tercero, f'RECAUDOS POR TERCERO - {proyecto}', subtitulo,
        [(g['movimientos'], f"Total {g['nit']} {g['nombre']}", g['totales']) for g in informe['terceros']],
        informe['totales'],
    )

    por_recibo = book.create_sheet('Por recibo')
    _hoja_agrupada(
        por_recibo, f'RECAUDOS POR RECIBO - {proyecto}', subtitulo,
        [(r['lineas'], f"Total recibo {r['recibo']}", r['totales']) for r in informe['recibos']],
        informe['totales'],
    )
    if vista == 'recibo':
        book.active = 1

    terceros = book.create_sheet('Terceros')
    _encabezados(terceros, 1, _COLS_TERCEROS)
    terceros.freeze_panes = 'A2'
    campos = ('tipo_doc', 'nit', 'nombre', 'direccion', 'ciudad', 'departamento', 'pais', 'email', 'celular', 'telefono')
    datos_terceros = informe.get('datos_terceros', [])
    for fila, t in enumerate(datos_terceros, start=2):
        for col, campo in enumerate(campos, start=1):
            terceros.cell(fila, col, t[campo])
    terceros.auto_filter.ref = f'A1:{get_column_letter(len(campos))}{len(datos_terceros) + 1}'

    resumen = book.create_sheet('Resumen contable')
    for col, (titulo, ancho) in enumerate((('Concepto', 32), ('Cuenta', 16), ('Recibos', 10),
                                           ('Débito', 18), ('Crédito', 18)), start=1):
        celda = resumen.cell(1, col, titulo)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = _ENCABEZADO
        resumen.column_dimensions[get_column_letter(col)].width = ancho
    fila = 2
    for r in informe['resumen']:
        resumen.cell(fila, 1, r['concepto'])
        resumen.cell(fila, 2, r['cuenta'])
        resumen.cell(fila, 3, r['recibos'])
        resumen.cell(fila, 4, float(r['debito'])).number_format = _NUM
        resumen.cell(fila, 5, float(r['credito'])).number_format = _NUM
        fila += 1
    resumen.cell(fila, 1, 'Sumas iguales')
    resumen.cell(fila, 4, float(informe['total_debito'])).number_format = _NUM
    resumen.cell(fila, 5, float(informe['total_credito'])).number_format = _NUM
    for col in range(1, 6):
        resumen.cell(fila, col).font = _NEGRITA
        resumen.cell(fila, col).border = Border(top=_LINEA)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
