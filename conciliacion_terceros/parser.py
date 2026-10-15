"""
Lector del Excel "Balance de prueba por tercero" exportado desde Alegra.

Estructura observada (oct-2026):
- Filas 1-8: título, empresa, "NIT.xxx", "Desde dd/mm/aaaa Hasta dd/mm/aaaa", cuenta filtrada, fecha de generación.
- Encabezado: Nivel | Código | Cuenta contable | Identificación | Nombre del tercero | Saldo anterior | Débito | Crédito | Saldo final
- Niveles: "> > > Cuenta", "> > > > Subcuenta", "> > > > > Total Auxiliar" (total por cuenta), "> > > > > Auxiliar" (tercero).
El saldo crédito viene positivo. Se valida que la suma de terceros cuadre con el "Total Auxiliar".
"""
import datetime
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl


class BalanceParseError(Exception):
    pass


@dataclass
class CuentaBalance:
    codigo: str
    nombre: str
    total: Decimal = None
    terceros: 'OrderedDict[str, dict]' = field(default_factory=OrderedDict)

    @property
    def suma(self):
        return sum((t['saldo'] for t in self.terceros.values()), Decimal('0'))


@dataclass
class BalanceTercero:
    empresa_nombre: str
    nit: str
    fecha_desde: datetime.date
    fecha_hasta: datetime.date
    cuentas: 'OrderedDict[str, CuentaBalance]'


def norm_ident(value):
    return re.sub(r'[^0-9A-Za-z]', '', str(value or '')).upper()


def _dec(value):
    if value in (None, ''):
        return Decimal('0')
    try:
        return Decimal(str(value)).quantize(Decimal('0.01'))
    except InvalidOperation:
        raise BalanceParseError(f'Valor numérico inválido: {value!r}')


def _parse_fecha(text):
    return datetime.datetime.strptime(text, '%d/%m/%Y').date()


def parse_balance(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise BalanceParseError(f'No se pudo abrir el archivo: {exc}')
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    empresa_nombre, nit, desde, hasta, header_idx = '', '', None, None, None
    for i, row in enumerate(rows[:30]):
        first = str(row[0] or '').strip() if row else ''
        if first.upper().startswith('NIT'):
            nit = norm_ident(first[3:])
        m = re.search(r'Desde\s+(\d{2}/\d{2}/\d{4})\s+Hasta\s+(\d{2}/\d{2}/\d{4})', first)
        if m:
            desde, hasta = _parse_fecha(m.group(1)), _parse_fecha(m.group(2))
        if first == 'Nivel':
            header_idx = i
            break
        if i == 1:
            empresa_nombre = first
    if header_idx is None:
        raise BalanceParseError('No se encontró el encabezado "Nivel | Código | ..." del balance por tercero.')
    if not hasta:
        raise BalanceParseError('No se encontró el rango "Desde ... Hasta ..." en el encabezado del archivo.')

    cuentas = OrderedDict()
    for row in rows[header_idx + 1:]:
        if not row or not row[0]:
            continue
        nivel = str(row[0]).strip()
        codigo = str(row[1] or '').strip()
        if not codigo:
            continue
        if nivel.endswith('Total Auxiliar'):
            cta = cuentas.setdefault(codigo, CuentaBalance(codigo, str(row[2] or '').strip()))
            cta.total = _dec(row[8])
        elif nivel.endswith('Auxiliar'):
            cta = cuentas.setdefault(codigo, CuentaBalance(codigo, str(row[2] or '').strip()))
            ident = norm_ident(row[3]) or f'SIN-ID-{len(cta.terceros) + 1}'
            linea = cta.terceros.get(ident)
            saldo = _dec(row[8])
            if linea:
                linea['saldo'] += saldo
            else:
                cta.terceros[ident] = {'identificacion': ident, 'nombre': str(row[4] or '').strip(), 'saldo': saldo}

    if not cuentas:
        raise BalanceParseError('El archivo no tiene filas de terceros ("Auxiliar").')
    for cta in cuentas.values():
        if cta.total is not None and abs(cta.suma - cta.total) > Decimal('1'):
            raise BalanceParseError(
                f'La cuenta {cta.codigo} no cuadra: suma de terceros {cta.suma:,.2f} vs total {cta.total:,.2f}. '
                f'¿El archivo está completo?'
            )
    return BalanceTercero(empresa_nombre, nit, desde, hasta, cuentas)


MESES = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12}


def _fecha_siigo(text):
    """'SEP/30/2026' -> date."""
    m = re.match(r'\s*([A-Za-z]{3})\s*/\s*(\d{1,2})\s*/\s*(\d{4})', str(text or ''))
    if not m or m.group(1).upper() not in MESES:
        return None
    return datetime.date(int(m.group(3)), MESES[m.group(1).upper()], int(m.group(2)))


def parse_siigo_terceros_cuenta(path):
    """
    Reporte SIIGO "TERCEROS DETALLADO POR CUENTA" (oct-2026, TERRANOVA):
    - A1 'Siigo - <EMPRESA>', J1 fecha de corte 'SEP/30/2026', A3 'De : ENE /2026 A : SEP /2026'.
    - Encabezado: CUENTA | NOMBRE | NIT | DIG.VER. | SUC | NOMBRE | SALDO ANTERIOR | DEBITOS | CREDITOS | NUEVO SALDO
    - Filas 'Total <cuenta>' y 'Total general'.
    SIIGO muestra el saldo crédito en negativo; para cuentas de naturaleza crédito (clases 2, 3 y 4)
    se invierte el signo para dejarlo igual que Alegra (crédito positivo).
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise BalanceParseError(f'No se pudo abrir el archivo: {exc}')
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    if not rows or not str(rows[0][0] or '').strip().lower().startswith('siigo'):
        raise BalanceParseError('No parece un reporte de SIIGO (la celda A1 debe empezar por "Siigo - ").')

    empresa_nombre = str(rows[0][0]).split('-', 1)[-1].strip()
    hasta = None
    for v in rows[0]:
        hasta = _fecha_siigo(v) or hasta
    header_idx = None
    for i, row in enumerate(rows[:20]):
        if row and str(row[0] or '').strip().upper() == 'CUENTA':
            header_idx = i
            break
    if header_idx is None:
        raise BalanceParseError('No se encontró el encabezado "CUENTA | NOMBRE | NIT | ..." del reporte.')
    if not hasta:
        raise BalanceParseError('No se encontró la fecha de corte (celda J1, p. ej. "SEP/30/2026").')
    desde = None
    m = re.search(r'De\s*:\s*([A-Za-z]{3})\s*/\s*(\d{4})', str(rows[2][0] or '') if len(rows) > 2 else '')
    if m and m.group(1).upper() in MESES:
        desde = datetime.date(int(m.group(2)), MESES[m.group(1).upper()], 1)

    cuentas = OrderedDict()
    for row in rows[header_idx + 1:]:
        if not row or row[0] in (None, ''):
            continue
        first = str(row[0]).strip()
        if first.lower().startswith('total'):
            codigo = first[5:].strip()
            if codigo in cuentas:
                cuentas[codigo].total = _dec(row[9]) * (-1 if codigo[:1] in '234' else 1)
            continue
        codigo = re.sub(r'\.0$', '', first)
        signo = -1 if codigo[:1] in '234' else 1
        cta = cuentas.setdefault(codigo, CuentaBalance(codigo, str(row[1] or '').strip()))
        ident = norm_ident(re.sub(r'\.0$', '', str(row[2] or ''))) or f'SIN-ID-{len(cta.terceros) + 1}'
        saldo = _dec(row[9]) * signo
        linea = cta.terceros.get(ident)
        if linea:
            linea['saldo'] += saldo
        else:
            cta.terceros[ident] = {'identificacion': ident, 'nombre': ' '.join(str(row[5] or '').split()), 'saldo': saldo}

    if not cuentas:
        raise BalanceParseError('El reporte no tiene filas de terceros.')
    for cta in cuentas.values():
        if cta.total is not None and abs(cta.suma - cta.total) > Decimal('1'):
            raise BalanceParseError(
                f'La cuenta {cta.codigo} no cuadra: suma de terceros {cta.suma:,.2f} vs total {cta.total:,.2f}. '
                f'¿El archivo está completo?'
            )
    return BalanceTercero(empresa_nombre, '', desde, hasta, cuentas)
