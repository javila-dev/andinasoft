from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from conciliacion_terceros.models import ConciliacionLinea

FMT_MONEY = '#,##0;[Red]-#,##0'
HEADER_FILL = PatternFill('solid', fgColor='1F3A34')
HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
THIN = Border(*(Side(style='thin', color='D5D8DC'),) * 4)
ESTADO_FILL = {
    ConciliacionLinea.ESTADO_CUADRADO: 'E8F5E9',
    ConciliacionLinea.ESTADO_DIFERENCIA: 'FFF8E1',
    ConciliacionLinea.ESTADO_SOLO_CONTABLE: 'FDECEA',
    ConciliacionLinea.ESTADO_SOLO_ANDINA: 'E3F2FD',
}


def _header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(ws.max_row, i)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, THIN
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(ws.max_row + 1, 1)
    ws.auto_filter.ref = f'A{ws.max_row}:{get_column_letter(len(cols))}{ws.max_row}'


def _row(ws, values, money_cols=(), fill=None):
    ws.append(values)
    r = ws.max_row
    for i in range(1, len(values) + 1):
        c = ws.cell(r, i)
        c.border = THIN
        c.font = Font(name='Calibri', size=10)
        if i in money_cols:
            c.number_format = FMT_MONEY
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)


def construir_libro(conciliacion):
    fuentes = list(conciliacion.fuentes.select_related('empresa'))
    resumen = conciliacion.resumen or {}
    lineas = list(conciliacion.lineas.select_related('revisado_por').order_by('identificacion'))
    estados = dict(ConciliacionLinea.ESTADO_CHOICES)
    revisiones = dict(ConciliacionLinea.REVISION_CHOICES)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = f'Conciliación de saldos de terceros · corte {conciliacion.fecha_corte:%d/%m/%Y}'
    ws['A1'].font = Font(name='Calibri', bold=True, size=14)
    ws['A2'] = f'Tolerancia ${conciliacion.tolerancia:,.0f} · ' + ' · '.join(f'{f.empresa.nombre}: {f.archivo_nombre}' for f in fuentes)
    ws.append([])
    _header(ws, ['Andinasoft · proyecto', 'Terceros', 'Valor'], [32, 12, 20])
    for p in resumen.get('proyectos', []):
        _row(ws, [p['proyecto'], p['terceros'], Decimal(p['valor'])], money_cols=(3,))
    ws.append([])
    _header(ws, ['Contabilidad · empresa', 'Cuenta', 'Nombre', 'Terceros', 'Saldo'], [32, 14, 40, 12, 20])
    for c in resumen.get('cuentas', []):
        _row(ws, [c['empresa_nombre'], c['cuenta'], c['cuenta_nombre'], c['terceros'], Decimal(c['saldo'])], money_cols=(5,))
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    # Una fila por tercero.
    sh = wb.create_sheet('Terceros')
    _header(sh, ['Identificación', 'Nombre', 'Andinasoft', 'Contabilidad', 'Diferencia', 'Estado', 'Proyectos', 'Empresas · cuentas',
                 'Observación', 'Alertas', 'Revisión', 'Nota', 'Revisado por'],
            [14, 34, 16, 16, 16, 16, 30, 50, 50, 60, 20, 40, 16])
    for l in sorted(lineas, key=lambda x: -abs(x.diferencia)):
        proyectos = ', '.join(sorted({f"{n['proyecto']} {n['ref']}" for n in l.negocios if n['suma']}))
        cuentas = ' | '.join(f"{c['empresa_nombre']} {c['cuenta']}: {Decimal(c['saldo']):,.0f}" for c in l.cuentas)
        _row(sh, [l.identificacion, l.nombre, l.saldo_andina, l.saldo_contable, l.diferencia, estados.get(l.estado, l.estado),
                  proyectos, cuentas, l.observacion, ' | '.join(a['texto'] for a in l.alertas), revisiones.get(l.revision, l.revision),
                  l.nota, (l.revisado_por.get_username() if l.revisado_por else '')],
             money_cols=(3, 4, 5), fill=ESTADO_FILL.get(l.estado))

    # Detalle débito / crédito: una fila por negocio y por saldo contable (sirve para tablas dinámicas).
    dt = wb.create_sheet('Detalle')
    _header(dt, ['Identificación', 'Nombre', 'Lado', 'Proyecto / empresa', 'Negocio / cuenta', 'Situación / nombre cuenta', 'Suma',
                 'Andinasoft', 'Contabilidad', 'Estado del tercero'], [14, 34, 13, 30, 22, 36, 8, 16, 16, 18])
    for l in lineas:
        for n in l.negocios:
            _row(dt, [l.identificacion, l.nombre, 'Andinasoft', n['proyecto'], n['ref'], n['situacion_label'], 'Sí' if n['suma'] else 'No',
                      Decimal(n['valor']) if n['suma'] else Decimal(0), None, estados.get(l.estado, l.estado)], money_cols=(8, 9))
        for c in l.cuentas:
            _row(dt, [l.identificacion, l.nombre, 'Contabilidad', c['empresa_nombre'], c['cuenta'], c['cuenta_nombre'], 'Sí',
                      None, Decimal(c['saldo']), estados.get(l.estado, l.estado)], money_cols=(8, 9))
    # Ajustes registrados (vigentes y anulados): base para generar los journals.
    aj = wb.create_sheet('Ajustes')
    _header(aj, ['Fecha', 'Grupo', 'Tipo', 'Origen tercero', 'Origen empresa', 'Origen cuenta', 'Destino tercero', 'Destino empresa',
                 'Destino cuenta', 'Valor', 'Entre empresas', 'Nota', 'Estado', 'Journal', 'Registrado por'],
            [16, 12, 16, 14, 26, 26, 14, 26, 26, 16, 12, 40, 12, 12, 16])
    nombres = {f.empresa_id: f.empresa.nombre for f in fuentes}
    for a in conciliacion.ajustes.select_related('created_by').order_by('created_at', 'id'):
        _row(aj, [a.created_at.strftime('%d/%m/%Y %H:%M'), a.grupo[:8], a.get_tipo_display(), a.origen_identificacion,
                  nombres.get(a.origen_empresa_id, a.origen_empresa_id), f'{a.origen_cuenta} {a.origen_cuenta_nombre}',
                  a.destino_identificacion, nombres.get(a.destino_empresa_id, a.destino_empresa_id), f'{a.destino_cuenta} {a.destino_cuenta_nombre}',
                  a.valor, 'No' if a.misma_empresa else 'Sí', a.nota, 'Anulado' if a.anulado else 'Vigente', a.get_journal_estado_display(),
                  a.created_by.get_username() if a.created_by else ''], money_cols=(10,), fill='F2F3F4' if a.anulado else None)
    return wb
